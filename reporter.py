import sqlite3
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NULL = ("N/A", "nan", "", "None")

NAVY         = "1A1A2E"
WHITE        = "FFFFFF"
BLUE_ACCENT  = "4361EE"
GREEN_FOUND  = "2D6A4F"
GREEN_LIGHT  = "D8F3DC"
YELLOW_LIGHT = "FFF3CD"
RED_LIGHT    = "FFEDE8"
BORDER_CLR   = "C5CAE9"
PURPLE       = "6A4C93"
GOLD         = "B8860B"


def _border():
    return Border(
        left=Side(style="thin", color=BORDER_CLR), right=Side(style="thin", color=BORDER_CLR),
        top=Side(style="thin", color=BORDER_CLR),  bottom=Side(style="thin", color=BORDER_CLR),
    )


def _autosize(ws, mn=10, mx=60):
    for col in ws.columns:
        max_len = 0
        letter  = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[letter].width = max(mn, min(max_len + 4, mx))


class Reporter:
    """Builds the client-ready Excel report and SQLite database."""

    def __init__(
        self,
        industry:    str,
        market:      str,
        report_file: str,
        db_file:     str,
        top_n:       int = 20,
    ):
        self.industry    = industry
        self.market      = market
        self.report_file = report_file
        self.db_file     = db_file
        self.top_n       = top_n

    # ---- Public API ---------------------------------------------------------

    def build_excel(self, df: pd.DataFrame) -> None:
        col_order = [
            "Ranking", "Company", "State", "Lead_Score", "Tier",
            "Sqm_2025", "Rank_2024", "Sqm_2024", "Growth_Rank", "Growth_Sqm_Pct",
            "B3_Ticker", "Age", "Capital", "Capital_Note", "Founded",
            "Official_Email", "Official_Phone", "Decision_Makers",
            "IG_Followers", "IG_Posts",
            "Website", "Email", "WhatsApp", "Instagram", "LinkedIn",
            "Growth_Story", "Score_Breakdown", "News_Highlights",
        ]
        col_order = [c for c in col_order if c in df.columns]
        df        = df[col_order].copy()

        top_cols = [
            "Ranking", "Company", "State", "Lead_Score", "Tier",
            "Growth_Rank", "Growth_Sqm_Pct", "B3_Ticker",
            "Age", "Capital", "Capital_Note", "Decision_Makers",
            "Official_Email", "Official_Phone", "WhatsApp",
            "IG_Followers", "Website", "Growth_Story", "News_Highlights",
        ]
        top_cols = [c for c in top_cols if c in df.columns]
        df_top   = df.head(self.top_n)[top_cols].copy()

        wb = Workbook()
        wb.remove(wb.active)
        ws_sum  = wb.create_sheet("📊 Summary")
        ws_top  = wb.create_sheet("🏆 Top Leads")
        ws_main = wb.create_sheet("📋 All Companies")

        self._build_summary_sheet(ws_sum, df)
        self._build_top_leads_sheet(ws_top, df_top)
        self._build_all_companies_sheet(ws_main, df)
        wb.active = wb["📊 Summary"]
        wb.save(self.report_file)
        print(f"  Excel report saved: {self.report_file}")

    def build_database(self, df: pd.DataFrame) -> None:
        conn = sqlite3.connect(self.db_file)
        cur  = conn.cursor()

        cur.execute("DROP TABLE IF EXISTS companies")
        cur.execute("""
            CREATE TABLE companies (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                ranking         INTEGER,
                company         TEXT,
                state           TEXT,
                sqm_2025        TEXT,
                website         TEXT,
                email           TEXT,
                whatsapp        TEXT,
                instagram       TEXT,
                linkedin        TEXT,
                official_email  TEXT,
                official_phone  TEXT,
                decision_makers TEXT,
                founded         TEXT,
                age             TEXT,
                capital         TEXT,
                capital_note    TEXT,
                b3_ticker       TEXT,
                growth_story    TEXT,
                news_highlights TEXT,
                lead_score      TEXT DEFAULT '0',
                tier            TEXT DEFAULT 'N/A'
            )
        """)

        for _, row in df.iterrows():
            cur.execute("""
                INSERT INTO companies (
                    ranking, company, state, sqm_2025,
                    website, email, whatsapp, instagram, linkedin,
                    official_email, official_phone, decision_makers,
                    founded, age, capital, capital_note,
                    b3_ticker, growth_story, news_highlights, lead_score, tier
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                str(row.get("Ranking",         "N/A")),
                str(row.get("Company",         "N/A")),
                str(row.get("State",           "N/A")),
                str(row.get("Sqm_2025",        "N/A")),
                str(row.get("Website",         "N/A")),
                str(row.get("Email",           "N/A")),
                str(row.get("WhatsApp",        "N/A")),
                str(row.get("Instagram",       "N/A")),
                str(row.get("LinkedIn",        "N/A")),
                str(row.get("Official_Email",  "N/A")),
                str(row.get("Official_Phone",  "N/A")),
                str(row.get("Decision_Makers", "N/A")),
                str(row.get("Founded",         "N/A")),
                str(row.get("Age",             "N/A")),
                str(row.get("Capital",         "N/A")),
                str(row.get("Capital_Note",    "N/A")),
                str(row.get("B3_Ticker",       "N/A")),
                str(row.get("Growth_Story",    "N/A")),
                str(row.get("News_Highlights", "N/A")),
                str(row.get("Lead_Score",      "0")),
                str(row.get("Tier",            "N/A")),
            ))

        conn.commit()
        conn.close()
        print(f"  Database saved: {self.db_file}")

    # ---- Sheet builders -----------------------------------------------------

    def _build_all_companies_sheet(self, ws, df: pd.DataFrame) -> None:
        headers = list(df.columns)
        ws.append(headers)
        self._style_header(ws)
        for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
            ws.append(list(row))
            tier  = str(row.get("Tier", ""))
            score = row.get("Lead_Score", 0)
            bg    = GREEN_LIGHT  if "Hot"  in tier else (
                    YELLOW_LIGHT if "Warm" in tier else (
                    RED_LIGHT    if row_idx % 2 == 0 else "FFF8F6"))
            for col_idx, cell in enumerate(
                ws.iter_rows(min_row=row_idx, max_row=row_idx).__next__(), start=1
            ):
                col_name = headers[col_idx - 1] if col_idx <= len(headers) else ""
                self._style_cell(cell, col_name, bg, score)
            ws.row_dimensions[row_idx].height = 18
        ws.freeze_panes = "A2"
        _autosize(ws)

    def _build_top_leads_sheet(self, ws, df_top: pd.DataFrame) -> None:
        last_col = get_column_letter(max(len(df_top.columns), 1))
        ws.merge_cells(f"A1:{last_col}1")
        t = ws["A1"]
        t.value     = f"🏆  TOP {len(df_top)} LEADS — CONSTRUTORAS SUDESTE 2025"
        t.font      = Font(bold=True, size=14, color=WHITE, name="Calibri")
        t.fill      = PatternFill("solid", fgColor=NAVY)
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells(f"A2:{last_col}2")
        s = ws["A2"]
        s.value     = f"Ranked by enriched Lead Score v3.0 · {datetime.now().strftime('%d %B %Y')}"
        s.font      = Font(italic=True, size=10, color="555555", name="Calibri")
        s.fill      = PatternFill("solid", fgColor="F0F4FF")
        s.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        ws.append([])
        ws.append(list(df_top.columns))
        self._style_header(ws, row_num=4)

        headers = list(df_top.columns)
        for pos, (_, row) in enumerate(df_top.iterrows(), start=1):
            ws.append(list(row))
            row_idx = pos + 4
            bg      = "FFFDE7" if pos == 1 else (
                      GREEN_LIGHT  if "Hot"  in str(row.get("Tier", "")) else
                      YELLOW_LIGHT if "Warm" in str(row.get("Tier", "")) else "FFF8F6")
            score   = row.get("Lead_Score", 0)
            for col_idx, cell in enumerate(
                ws.iter_rows(min_row=row_idx, max_row=row_idx).__next__(), start=1
            ):
                col_name = headers[col_idx - 1] if col_idx <= len(headers) else ""
                self._style_cell(cell, col_name, bg, score)
                if col_name == "Lead_Score" and pos == 1:
                    cell.font = Font(bold=True, size=13, color="F57F17", name="Calibri")
            ws.row_dimensions[row_idx].height = 36
        ws.freeze_panes = "A5"
        _autosize(ws)

    def _build_summary_sheet(self, ws, df: pd.DataFrame) -> None:
        hot   = len(df[df["Tier"] == "🔥 Hot Lead"])
        warm  = len(df[df["Tier"] == "⚡ Warm Lead"])
        cold  = len(df[df["Tier"] == "❄️ Cold Lead"])
        total = len(df)
        avg   = round(df["Lead_Score"].mean(), 1)
        top   = df.iloc[0]

        def pct(col):
            if col not in df.columns:
                return "N/A"
            n = df[col].apply(lambda x: str(x) not in NULL).sum()
            return f"{n} / {total}  ({round(n / total * 100)}%)"

        b3_count = (
            df["B3_Ticker"].apply(lambda x: str(x) not in ("Private", "N/A", "nan", "")).sum()
            if "B3_Ticker" in df.columns else 0
        )
        climbers = (
            df[df["Growth_Rank"].apply(
                lambda x: str(x) not in NULL and not str(x).startswith("-") and str(x) != "0"
            )] if "Growth_Rank" in df.columns else pd.DataFrame()
        )
        top_climber = climbers.iloc[0] if len(climbers) > 0 else None

        ws.merge_cells("A1:C1")
        t = ws["A1"]
        t.value     = f"EXECUTIVE SUMMARY — {self.industry.upper()} {self.market.upper()} 2025"
        t.font      = Font(bold=True, size=13, color=WHITE, name="Calibri")
        t.fill      = PatternFill("solid", fgColor=NAVY)
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        rows_data = [
            ("", ""), ("📊 DATASET", ""),
            ("Total Companies", total),
            ("Market",          self.market),
            ("Source",          "INTEC 2025 + INTEC 2024 + Receita Federal API"),
            ("Generated",       datetime.now().strftime("%d %B %Y %H:%M")),
            ("", ""), ("🏆 LEAD TIERS", ""),
            ("🔥 Hot Leads (65–100)", hot),
            ("⚡ Warm Leads (45–64)", warm),
            ("❄️ Cold Leads (0–44)",  cold),
            ("Average Score",         avg),
            ("", ""), ("📈 GROWTH INTELLIGENCE", ""),
            ("Top Scoring Company",
             f"{top['Company']}  ({top['Lead_Score']} pts)"),
        ]
        if top_climber is not None:
            rows_data.append((
                "Biggest Ranking Climber",
                f"{top_climber['Company']}  (+{top_climber['Growth_Rank']} positions)",
            ))
        rows_data += [
            ("", ""), ("📞 CONTACT COVERAGE", ""),
            ("Official Emails (Receita Federal)", pct("Official_Email")),
            ("Website Emails",                    pct("Email")),
            ("Official Phones",                   pct("Official_Phone")),
            ("WhatsApp",                          pct("WhatsApp")),
            ("Instagram",                         pct("Instagram")),
            ("LinkedIn",                          pct("LinkedIn")),
            ("Decision Makers",                   pct("Decision_Makers")),
            ("B3 Listed Companies",               f"{b3_count} / {total}"),
            ("", ""), ("💡 SCORING METHODOLOGY v3.0", ""),
            ("National Position",       "Up to 20 pts"),
            ("Construction Volume m²",  "Up to 20 pts  ← primary signal"),
            ("YoY Ranking Growth",      "Up to 20 pts"),
            ("m² Volume Growth",        "Up to 10 pts"),
            ("Contact Accessibility",   "Up to 20 pts  (official channels weighted higher)"),
            ("Capital & Maturity",      "Up to 10 pts  (B3 status + age + verified capital)"),
            ("Maximum Score",           "100 pts"),
            ("🔥 Hot threshold",        "65+ pts"),
            ("⚡ Warm threshold",       "45–64 pts"),
            ("❄️ Cold threshold",       "0–44 pts"),
            ("", ""), ("📌 DATA SOURCES", ""),
            ("Ranking data",   "INTEC 2025 + INTEC 2024"),
            ("Contact data",   "Company websites (Playwright)"),
            ("Official data",  "Receita Federal via ReceitaWS API"),
            ("Social media",   "Instagram public profiles"),
        ]

        for r_idx, (label, value) in enumerate(rows_data, start=2):
            lc = ws.cell(row=r_idx, column=1, value=label)
            vc = ws.cell(row=r_idx, column=2, value=value)
            is_section = any(label.startswith(e) for e in ["📊", "🏆", "📈", "📞", "💡", "📌"])
            if is_section:
                for c in (lc, vc):
                    c.font = Font(bold=True, size=11, color=NAVY, name="Calibri")
                    c.fill = PatternFill("solid", fgColor="E8ECFF")
                ws.row_dimensions[r_idx].height = 22
            elif label:
                lc.font = Font(size=10, color="333333", name="Calibri")
                vc.font = Font(size=10, bold=True, color=NAVY, name="Calibri")
                for c in (lc, vc):
                    c.fill = PatternFill("solid", fgColor="F8F9FF")
                ws.row_dimensions[r_idx].height = 18
            for c in (lc, vc):
                c.border    = _border()
                c.alignment = Alignment(horizontal="left", vertical="center")

        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 50

    # ---- Styling helpers ----------------------------------------------------

    def _style_header(self, ws, row_num: int = 1) -> None:
        for cell in ws[row_num]:
            cell.fill      = PatternFill("solid", fgColor=NAVY)
            cell.font      = Font(bold=True, color=WHITE, size=11, name="Calibri")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = _border()
        ws.row_dimensions[row_num].height = 28

    def _style_cell(self, cell, col_name: str, bg: str, score=0) -> None:
        cell.fill      = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _border()
        cell.font      = Font(size=10, name="Calibri")

        if col_name == "Ranking":
            cell.font = Font(bold=True, color=BLUE_ACCENT, size=10, name="Calibri")
        elif col_name == "Lead_Score":
            color = "1B5E20" if score >= 65 else "E65100" if score >= 45 else "B71C1C"
            cell.font = Font(bold=True, size=11, color=color, name="Calibri")
        elif col_name in (
            "Email", "WhatsApp", "Instagram", "LinkedIn",
            "Official_Email", "Official_Phone",
        ):
            v = str(cell.value) if cell.value else ""
            if v and v not in NULL:
                cell.font = Font(color=GREEN_FOUND, size=10, name="Calibri")
        elif col_name == "Decision_Makers":
            v = str(cell.value) if cell.value else ""
            if v and v not in NULL:
                cell.font = Font(color=PURPLE, size=10, name="Calibri")
        elif col_name == "B3_Ticker":
            v = str(cell.value) if cell.value else ""
            if v and v not in ("Private",) + NULL:
                cell.font = Font(bold=True, color=GOLD, size=10, name="Calibri")
        elif col_name == "Capital_Note":
            v = str(cell.value) if cell.value else ""
            if v.startswith("✅"):
                cell.font = Font(color=GREEN_FOUND, size=9, name="Calibri")
            elif v.startswith("⚠️"):
                cell.font = Font(color="E65100", size=9, name="Calibri")
            elif v.startswith("❌"):
                cell.font = Font(color="B71C1C", size=9, name="Calibri")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        elif col_name == "Growth_Story":
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.font      = Font(size=9, color="444444", name="Calibri", italic=True)
        elif col_name == "News_Highlights":
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            cell.font      = Font(size=9, color="1A3A5C", name="Calibri", italic=True)

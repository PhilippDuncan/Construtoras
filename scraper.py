import asyncio
import re
import logging
from urllib.parse import urlparse

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeout

NULL = ("N/A", "nan", "", "None")

_SKIP_DOMAINS = [
    "google.", "youtube.", "facebook.", "instagram.", "linkedin.",
    "wikipedia.", "jusbrasil.", "apto.vc", "imovelweb.", "vivareal.",
    "olx.", "zap.", "twitter.", "tiktok.", "duckduckgo.", "bing.",
    "chavesnamao.", "99lotes.", "lopes.", "creci.", "reclameaqui.",
    "infomoney.", "exame.", "estadao.", "folha.", "uol.", "globo.",
    "construtoras.net", "guiadaconstrucao", "reddit.", "cnpj.biz",
    "cnpj.info", "gupy.io", "incorporacaoimobiliaria.", "hubimobiliario.",
    "sienge.com.br", "seudinheiro.", "bbc.com",
]
_GENERIC_WORDS = {
    "construtora", "engenharia", "incorporadora", "empreendimentos",
    "construções", "construcoes", "ltda", "grupo", "inc", "e", "de",
    "do", "da", "dos", "das", "em", "incorporação",
}


class Scraper:
    """Scrapes INTEC ranking + company websites, applies manual corrections."""

    _USER_AGENT = (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    )

    def __init__(self, cfg):
        self.cfg = cfg
        logging.basicConfig(
            level=logging.INFO,
            format="%(asctime)s | %(message)s",
            datefmt="%H:%M:%S",
        )
        self._log = logging.getLogger(__name__)

    async def run(self) -> pd.DataFrame:
        async with async_playwright() as p:
            browser = await p.chromium.launch(
                headless=True,
                args=["--disable-dev-shm-usage", "--no-sandbox", "--disable-gpu"],
            )
            companies = await self._run_pipeline(browser)
            await browser.close()

        df = pd.DataFrame(companies)
        df.sort_values("Ranking", inplace=True)
        df.reset_index(drop=True, inplace=True)
        df = self._apply_overrides(df)
        self._export_excel(df)
        print(f"  Saved {len(df)} companies to {self.cfg.RAW_FILE}")
        return df

    # ---- Internal orchestration ---------------------------------------------

    async def _run_pipeline(self, browser) -> list[dict]:
        context = await browser.new_context(user_agent=self._USER_AGENT)
        page    = await context.new_page()
        companies = await self._scrape_ranking(page)
        await context.close()

        await self._scrape_contacts(companies, browser)
        return companies

    # ---- Step 1: INTEC ranking ----------------------------------------------

    async def _scrape_ranking(self, page) -> list[dict]:
        for attempt in range(1, self.cfg.MAX_RETRIES + 1):
            try:
                await page.goto(
                    self.cfg.INTEC_URL_2025,
                    timeout=self.cfg.PAGE_TIMEOUT,
                    wait_until="domcontentloaded",
                )
                await page.wait_for_selector("table", timeout=self.cfg.PAGE_TIMEOUT)
                break
            except Exception as e:
                if attempt == self.cfg.MAX_RETRIES:
                    raise
                self._log.warning(f"Ranking attempt {attempt} failed: {e}")

        rows      = await page.locator("table tr").all()
        companies = []
        seen      = set()

        for row in rows:
            cells = await row.locator("td").all_inner_texts()
            if len(cells) < 5:
                continue
            pos    = cells[0].strip().replace("º", "").replace("°", "")
            name   = cells[1].strip().upper()
            sqm    = cells[2].strip()
            region = cells[3].strip().upper()
            state  = cells[4].strip().upper()
            if not pos.isdigit() or not name:
                continue
            if self.cfg.TARGET_REGION not in region:
                continue
            if name in seen:
                continue
            seen.add(name)
            companies.append({
                "Ranking": int(pos),
                "Company": name,
                "Sqm_2025": sqm,
                "Region":  region,
                "State":   state,
            })

        print(f"  Ranking: {len(companies)} {self.cfg.TARGET_REGION} companies extracted.")
        return companies

    # ---- Step 2: website search + contact extraction ------------------------

    def _find_website(self, company_name: str) -> str | None:
        from ddgs import DDGS

        keywords = [
            w.lower() for w in company_name.split()
            if w.lower() not in _GENERIC_WORDS and len(w) > 2
        ]
        queries = [
            f'"{company_name}" site oficial construtora Brasil',
            f"{company_name} construtora incorporadora Brasil",
        ]
        for query in queries:
            try:
                with DDGS() as ddgs:
                    results = list(ddgs.text(query, max_results=8))
                br_candidates  = []
                any_candidates = []
                for r in results:
                    url   = r.get("href", "")
                    title = r.get("title", "").lower()
                    body  = r.get("body",  "").lower()
                    if not url:
                        continue
                    url_lower = url.lower()
                    if any(s in url_lower for s in _SKIP_DOMAINS):
                        continue
                    if keywords and not any(
                        kw in url_lower or kw in title or kw in body
                        for kw in keywords
                    ):
                        continue
                    parsed = urlparse(url)
                    root   = f"{parsed.scheme}://{parsed.netloc}/"
                    if any(tld in url_lower for tld in [".com.br", ".eng.br", ".net.br", ".org.br"]):
                        br_candidates.append(root)
                    else:
                        any_candidates.append(root)
                best = br_candidates[0] if br_candidates else (
                       any_candidates[0] if any_candidates else None)
                if best:
                    return best
            except Exception as e:
                self._log.warning(f"Search failed for '{company_name}': {e}")
        return None

    async def _extract_contacts(self, page, url: str) -> dict:
        result = {
            "Website": url, "Email": "N/A",
            "WhatsApp": "N/A", "Instagram": "N/A", "LinkedIn": "N/A",
        }
        try:
            await page.goto(url, timeout=self.cfg.VISIT_TIMEOUT, wait_until="domcontentloaded")
            await page.mouse.wheel(0, 2000)
            await asyncio.sleep(1)
            await page.mouse.wheel(0, 5000)
            await asyncio.sleep(1)
            result["Website"] = page.url

            html       = await page.content()
            raw_emails = re.findall(
                r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", html
            )
            result["Email"] = self._pick_best_email(raw_emails)

            links = await page.locator("a[href]").all()
            for link in links:
                try:
                    href = await link.get_attribute("href")
                    if not href:
                        continue
                    hl = href.lower()
                    if result["WhatsApp"] == "N/A" and any(
                        x in hl for x in ["wa.me", "whatsapp.com/send", "api.whatsapp"]
                    ):
                        result["WhatsApp"] = href
                    if result["Instagram"] == "N/A" and (
                        "instagram.com/" in hl and "instagram.com/p/" not in hl
                    ):
                        result["Instagram"] = href.split("?")[0].rstrip("/")
                    if result["LinkedIn"] == "N/A" and "linkedin.com/company/" in hl:
                        result["LinkedIn"] = href.split("?")[0].rstrip("/")
                except Exception:
                    continue

            if result["Email"] == "N/A":
                for link in links:
                    try:
                        href = await link.get_attribute("href")
                        if href and href.lower().startswith("mailto:"):
                            result["Email"] = href.replace("mailto:", "").split("?")[0].strip()
                            break
                    except Exception:
                        continue

        except PlaywrightTimeout:
            result["Website"] = url + " [TIMEOUT]"
        except asyncio.CancelledError:
            result["Website"] = url + " [CANCELLED]"
            raise
        except Exception:
            result["Website"] = url + " [ERROR]"
        return result

    async def _scrape_contacts(self, companies: list[dict], browser) -> None:
        context = None
        page    = None

        async def fresh_page():
            nonlocal context, page
            if context:
                try:
                    await context.close()
                except Exception:
                    pass
            context = await browser.new_context(user_agent=self._USER_AGENT)
            page    = await context.new_page()
            return page

        page  = await fresh_page()
        total = len(companies)

        for i, company in enumerate(companies, start=1):
            name = company["Company"]

            # Manual URL corrections take priority over DuckDuckGo search
            if name in self.cfg.MANUAL_URL_CORRECTIONS:
                website_url = self.cfg.MANUAL_URL_CORRECTIONS[name]
                print(f"  [{i:02d}/{total}] {name[:40]:<40} [MANUAL]", flush=True)
            else:
                if i > 1 and (i - 1) % self.cfg.CONTEXT_RESET == 0:
                    page = await fresh_page()
                print(f"  [{i:02d}/{total}] {name[:40]:<40}", end="", flush=True)
                website_url = None
                for attempt in range(1, self.cfg.MAX_RETRIES + 1):
                    try:
                        website_url = self._find_website(name)
                        break
                    except Exception as e:
                        if attempt == self.cfg.MAX_RETRIES:
                            self._log.warning(f"Website search failed: {name}")

            if not website_url:
                print(" NOT FOUND")
                company.update({
                    "Website": "N/A", "Email": "N/A",
                    "WhatsApp": "N/A", "Instagram": "N/A", "LinkedIn": "N/A",
                })
                continue

            try:
                contacts = await self._extract_contacts(page, website_url)
            except BaseException:
                page = await fresh_page()
                contacts = {
                    "Website": website_url + " [CRASHED]", "Email": "N/A",
                    "WhatsApp": "N/A", "Instagram": "N/A", "LinkedIn": "N/A",
                }

            company.update(contacts)
            e  = "✓" if contacts["Email"]     != "N/A" else "✗"
            wa = "✓" if contacts["WhatsApp"]  != "N/A" else "✗"
            ig = "✓" if contacts["Instagram"] != "N/A" else "✗"
            print(f" E:{e} WA:{wa} IG:{ig}")
            await asyncio.sleep(self.cfg.DELAY_BETWEEN)

        if context:
            try:
                await context.close()
            except Exception:
                pass

    # ---- Post-processing ----------------------------------------------------

    def _pick_best_email(self, emails: list[str]) -> str:
        bad_ext = (".png", ".jpg", ".jpeg", ".webp", ".svg", ".gif", ".pdf")
        bad_kw  = ("sentry", "example", "domain", "email@", "test@", "noreply")
        valid   = [
            e for e in emails
            if not any(e.lower().endswith(x) for x in bad_ext)
            and not any(k in e.lower() for k in bad_kw)
        ]
        if not valid:
            return "N/A"
        for kw in ["contato", "comercial", "vendas", "atendimento", "info"]:
            for em in valid:
                if kw in em.lower():
                    return em
        return list(dict.fromkeys(valid))[0]

    def _apply_overrides(self, df: pd.DataFrame) -> pd.DataFrame:
        for idx, row in df.iterrows():
            company = str(row["Company"])
            if company in self.cfg.EMAIL_OVERRIDES:
                df.at[idx, "Email"] = self.cfg.EMAIL_OVERRIDES[company]
            if company in self.cfg.INSTAGRAM_OVERRIDES:
                df.at[idx, "Instagram"] = self.cfg.INSTAGRAM_OVERRIDES[company]
        return df

    # ---- Excel export -------------------------------------------------------

    def _export_excel(self, df: pd.DataFrame) -> None:
        col_order = [
            "Ranking", "Company", "State", "Sqm_2025",
            "Website", "Email", "WhatsApp", "Instagram", "LinkedIn",
        ]
        col_order = [c for c in col_order if c in df.columns]
        df[col_order].to_excel(
            self.cfg.RAW_FILE, index=False, sheet_name=self.cfg.RAW_SHEET
        )
        self._format_excel(self.cfg.RAW_FILE)

    def _format_excel(self, path: str) -> None:
        wb      = load_workbook(path)
        ws      = wb.active
        NAVY    = "1A1A2E"; WHITE   = "FFFFFF"
        LAVENDER = "EEF2FF"; BLUE   = "4361EE"
        GREEN   = "2D6A4F"; BORDER = "C5CAE9"
        thin = Border(
            left=Side(style="thin", color=BORDER), right=Side(style="thin", color=BORDER),
            top=Side(style="thin", color=BORDER),  bottom=Side(style="thin", color=BORDER),
        )
        headers = [cell.value for cell in ws[1]]
        for cell in ws[1]:
            cell.fill      = PatternFill("solid", fgColor=NAVY)
            cell.font      = Font(bold=True, color=WHITE, size=11, name="Calibri")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = thin
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            bg = LAVENDER if row_idx % 2 == 0 else WHITE
            for col_idx, cell in enumerate(row, start=1):
                col_name = headers[col_idx - 1] if col_idx <= len(headers) else ""
                cell.fill      = PatternFill("solid", fgColor=bg)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border    = thin
                cell.font      = Font(size=10, name="Calibri")
                if col_name == "Ranking":
                    cell.font = Font(bold=True, color=BLUE, size=10, name="Calibri")
                if col_name in ("Email", "WhatsApp", "Instagram", "LinkedIn"):
                    v = str(cell.value) if cell.value else ""
                    if v and v not in NULL:
                        cell.font = Font(color=GREEN, size=10, name="Calibri")
        for col in ws.columns:
            max_len = 0
            letter  = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[letter].width = min(max_len + 4, 55)
        ws.freeze_panes = "A2"
        wb.save(path)
